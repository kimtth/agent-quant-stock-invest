"""Dump generic workbook sheets as text; output goes to CONFIG.output or stdout.

Workbook text is data, retained in its original language. Formula text is kept
instead of silently dropping formulas lacking cached results.
"""

from __future__ import annotations

from pathlib import Path


def workbook_text(path: str | Path, only: str | None = None) -> str:
    """Read a workbook explicitly; reject missing sheets and empty input."""
    from openpyxl import load_workbook

    source = Path(path)
    if not source.is_file():
        raise ValueError(f"Workbook not found: {source}; set workbook to an existing XLSX file")
    try:
        wb = load_workbook(source, data_only=False, read_only=True)
    except Exception as exc:
        raise ValueError(f"Cannot read workbook {source}; provide a valid XLSX workbook: {exc}") from exc
    try:
        if only is not None and only not in wb.sheetnames:
            raise ValueError(f"Sheet {only!r} not found; available sheets: {wb.sheetnames}")
        lines = ["SHEETS: " + ", ".join(wb.sheetnames)]
        nonempty = False
        for name in wb.sheetnames:
            if only is not None and only != name:
                continue
            lines.append(f"\n===== SHEET: {name} =====")
            for row in wb[name].iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    nonempty = True
                    lines.append(" | ".join(cells))
        if not nonempty:
            raise ValueError("Workbook selection is empty; supply at least one non-empty sheet")
        return "\n".join(lines) + "\n"
    except Exception as exc:
        raise ValueError(f"Cannot extract workbook {source}: {exc}") from exc
    finally:
        wb.close()


def main() -> None:
    from market_config import CONFIG
    if CONFIG.workbook is None:
        raise ValueError("Set workbook in the research config to an existing XLSX file")
    if CONFIG.output is not None and Path(CONFIG.output).resolve() == Path(CONFIG.workbook).resolve():
        raise ValueError("output must differ from workbook; choose a text output path")
    text = workbook_text(CONFIG.workbook)
    if CONFIG.output is None:
        print(text, end="")
    else:
        output = Path(CONFIG.output)
        try:
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(text, encoding="utf-8")
        except OSError as exc:
            raise ValueError(f"Cannot write output {output}; choose a writable text path: {exc}") from exc
        print(f"Workbook text written to {output}")


if __name__ == "__main__":
    main()
