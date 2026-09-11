"""Offline portfolio CSV, benchmark, workbook and import-safety regressions."""

from __future__ import annotations

import builtins
import importlib
import math
import sys
from dataclasses import replace
from pathlib import Path
from types import ModuleType

import pandas as pd
import pytest


TOOLS = Path(__file__).resolve().parents[1] / ".old" / "tools"
MODULES = (
    "portfolio_history", "analyze_drawdown", "analyze_drawdown_bounded",
    "analyze_drawdown_filters", "analyze_drawdown_frozen", "analyze_drawdown_liquidate",
    "analyze_drawdown_lockout", "compare_v1_v2", "list_out_of_market",
    "search_buy_block", "search_max_profit", "why_count_escape", "dump_factors",
    "dump_genport_guide",
)


@pytest.fixture
def modules(monkeypatch):
    monkeypatch.syspath_prepend(str(TOOLS))
    monkeypatch.delenv("GLOBAL_RESEARCH_CONFIG", raising=False)
    for name in (*MODULES, "market_config"):
        monkeypatch.delitem(sys.modules, name, raising=False)
    config = importlib.import_module("market_config")
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, start="2020-01-01", end="2021-01-01"))
    loaded = {name: importlib.import_module(name) for name in MODULES}
    loaded["market_config"] = config
    yield loaded
    # Remove imports made by this fixture before monkeypatch restores prior modules.
    for name in (*MODULES, "market_config"):
        sys.modules.pop(name, None)


def write_csv(tmp_path, body, header="date,return,asset,cash,held", name="history.csv"):
    path = tmp_path / name
    path.write_text(header + "\n" + body, encoding="utf-8-sig")
    return path


def test_english_and_original_aliases(modules, tmp_path):
    history = modules["portfolio_history"]
    english = write_csv(tmp_path, "2020-01-02,0.01,101,1,1,2\n", "date,return,asset,cash,held,buys")
    korean = write_csv(tmp_path, "2020-01-02,1,101,1,1,2\n", "날짜,일일수익률,총자산,남은 현금,남은 종목수,매수 종목수", "original.csv")
    assert history.read_history(english) == history.read_history(korean)
    assert history.read_history(english)[0].ret == 0.01
    assert history.read_history(english)[0].buys == 2
    plural = write_csv(tmp_path, "2020-01-02,1,101,1,1\n", "date,return,assets,cash,held", "plural.csv")
    assert history.read_history(plural)[0].ret == 1.0
    assert history.read_history(plural)[0].buys is None
    assert modules["analyze_drawdown"].Row is history.Row


@pytest.mark.parametrize("body, message", [
    ("", "no data rows"),
    ("bad,0,100,0,1\n", "date"),
    ("2020-02-30,0,100,0,1\n", "date"),
    ("2020-01-02,nan,100,0,1\n", "finite"),
    ("2020-01-02,inf,100,0,1\n", "finite"),
    ("2020-01-02,-1,100,0,1\n", "greater than -1"),
    ("2020-01-02,0,0,0,1\n", "positive"),
    ("2020-01-02,0,inf,0,1\n", "finite"),
    ("2020-01-02,0,100,-1,1\n", "cash"),
    ("2020-01-02,0,100,101,1\n", "cash"),
    ("2020-01-02,0,100,nan,1\n", "finite"),
    ("2020-01-02,0,100,0,-1\n", "integer"),
    ("2020-01-02,0,100,0,1.5\n", "integer"),
    ("2020-01-02,0,100,0,1\n2020-01-02,0,100,0,1\n", "unique"),
    ("2020-01-03,0,100,0,1\n2020-01-02,0,100,0,1\n", "increasing"),
    ("2020-01-02,0,100,0\n", "width"),
    ("2020-01-02,0,100,0,1,extra\n", "width"),
])
def test_rejects_malformed_history(modules, tmp_path, body, message):
    path = write_csv(tmp_path, body)
    with pytest.raises(ValueError, match=message):
        modules["portfolio_history"].read_history(path)


def test_required_headers_buys_and_date_range(modules, tmp_path):
    read = modules["portfolio_history"].read_history
    missing = write_csv(tmp_path, "2020-01-02,0\n", "date,return")
    with pytest.raises(ValueError, match="missing required column asset"):
        read(missing)
    path = write_csv(tmp_path, "2020-01-02,0,100,0,1\n2020-01-03,0.1,110,0,1\n")
    with pytest.raises(ValueError, match="buys"):
        read(path, require_buys=True)
    assert len(read(path, start="2020-01-02", end="2020-01-03")) == 1
    with pytest.raises(ValueError, match="no rows"):
        read(path, start="2020-02-01")
    with pytest.raises(ValueError, match="not found"):
        read(tmp_path / "absent.csv")
    for count in ("-1", "1.5", "nan"):
        bad = write_csv(tmp_path, f"2020-01-02,0,100,0,1,{count}\n", "date,return,asset,cash,held,buys")
        with pytest.raises(ValueError, match="buys"):
            read(bad)


def test_alias_ambiguity_and_full_validation(modules, tmp_path):
    read = modules["portfolio_history"].read_history
    path = write_csv(tmp_path, "2020-01-02,0,100,100,0,1\n", "date,return,asset,assets,cash,held")
    with pytest.raises(ValueError, match="ambiguous"):
        read(path)
    path = write_csv(tmp_path, "2020-01-02,0,100,100,0,1\n", "date,return,asset,asset,cash,held")
    with pytest.raises(ValueError, match="duplicate"):
        read(path)
    path = write_csv(tmp_path, "2019-12-31,nan,100,0,1\n2020-01-02,0,100,0,1\n")
    with pytest.raises(ValueError, match="finite"):
        read(path, start="2020-01-01")


@pytest.mark.parametrize("ticker", ["SPY", "QQQ"])
def test_benchmark_fallback_offline(modules, monkeypatch, capsys, ticker):
    config = modules["market_config"]
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, benchmark=ticker))
    market = ModuleType("market_data")
    calls = []

    def download(symbols, start, end):
        calls.append((symbols, start, end))
        return {ticker: pd.DataFrame({"Close": [100, 110, 99]}, index=pd.to_datetime(["2020-01-02", "2020-01-03", "2020-01-06"]))}

    market.download_ohlcv = download
    monkeypatch.setitem(sys.modules, "market_data", market)
    rows = modules["analyze_drawdown"].load()
    assert calls == [([ticker], "2020-01-01", "2021-01-01")]
    assert [r.ret for r in rows] == pytest.approx([0, 0.1, -0.1])
    assert [r.asset for r in rows] == pytest.approx([1, 1.1, 0.99])
    assert all(r.cash == 0 and r.held == 1 and r.buys is None for r in rows)
    output = capsys.readouterr().out
    assert ticker in output and "buy-and-hold" in output
    assert "Counterfactual" in output and "not out-of-sample" in output


def test_explicit_history_precedence(modules, monkeypatch, tmp_path):
    config = modules["market_config"]
    path = write_csv(tmp_path, "2020-01-02,0,100,0,1\n")
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, history_csv=tmp_path / "missing.csv"))
    assert len(modules["analyze_drawdown"].load(path)) == 1
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, history_csv=path))
    assert len(modules["analyze_drawdown"].load()) == 1


@pytest.mark.parametrize("dates, prices, message", [
    (["2020-01-02"], [100], "at least two"),
    (["2020-01-02", "2020-01-02"], [100, 101], "unique"),
    (["2020-01-03", "2020-01-02"], [100, 101], "increasing"),
    (["bad", "2020-01-02"], [100, 101], "dates"),
    (["2020-01-02", "2020-01-03"], [100, 0], "positive"),
    (["2020-01-02", "2020-01-03"], [100, float("nan")], "finite"),
])
def test_invalid_benchmark_data(modules, monkeypatch, dates, prices, message):
    market = ModuleType("market_data")
    market.download_ohlcv = lambda *args: {"SPY": pd.DataFrame({"Close": prices}, index=dates)}
    monkeypatch.setitem(sys.modules, "market_data", market)
    with pytest.raises(ValueError, match=message):
        modules["portfolio_history"].benchmark_history()


def test_comparison_requires_real_buys_and_matching_dates(modules, monkeypatch, tmp_path):
    compare = modules["compare_v1_v2"]
    config = modules["market_config"]
    with pytest.raises(ValueError, match="history_csv and compare_csv"):
        compare.comparison_histories()
    first = write_csv(tmp_path, "2020-01-02,0,100,0,1\n")
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, history_csv=first, compare_csv=first))
    with pytest.raises(ValueError, match="buys"):
        compare.comparison_histories()
    first = write_csv(tmp_path, "2020-01-02,0,100,0,1,2\n", "date,return,asset,cash,held,buys")
    assert compare.buys(first) == {"2020-01-02": 2}
    a, b = compare.comparison_histories()
    assert a == b
    second = write_csv(tmp_path, "2020-01-03,0,100,0,1,0\n", "date,return,asset,cash,held,buys", "second.csv")
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, compare_csv=second))
    with pytest.raises(ValueError, match="identical dates"):
        compare.comparison_histories()


def test_workbooks_generic_output_and_empty_errors(modules, monkeypatch, tmp_path, capsys):
    from openpyxl import Workbook
    guide = modules["dump_genport_guide"]
    config = modules["market_config"]
    with pytest.raises(ValueError, match="workbook"):
        guide.main()
    book = Workbook()
    book.active.title = "Overview"
    book.active.append(["Factor", "Description"])
    book.active.append(["Momentum", "원문 데이터"])
    book.create_sheet("Extra").append(["Formula", "=1+2"])
    path = tmp_path / "guide.xlsx"
    book.save(path)
    book.close()
    text = guide.workbook_text(path)
    assert "Overview" in text and "Extra" in text and "원문 데이터" in text and "=1+2" in text
    with pytest.raises(ValueError, match="not found"):
        guide.workbook_text(path, "Missing")
    output = tmp_path / "nested" / "guide.txt"
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, workbook=path, output=output))
    modules["dump_factors"].main()
    assert output.read_text(encoding="utf-8") == text
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, output=None))
    guide.main()
    assert "Momentum" in capsys.readouterr().out
    monkeypatch.setattr(config, "CONFIG", replace(config.CONFIG, output=path))
    with pytest.raises(ValueError, match="differ"):
        guide.main()
    empty = Workbook()
    empty_path = tmp_path / "empty.xlsx"
    empty.save(empty_path)
    empty.close()
    with pytest.raises(ValueError, match="empty"):
        guide.workbook_text(empty_path)
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(ValueError, match="valid XLSX"):
        guide.workbook_text(invalid)


def test_imports_have_no_history_workbook_network_or_grid_work(modules, monkeypatch, capsys):
    def forbidden(*args, **kwargs):
        raise AssertionError("Import must not perform data I/O or grid execution")

    # Config is preloaded. Target imports must only define functions/constants.
    market = ModuleType("market_data")
    market.download_ohlcv = forbidden
    monkeypatch.setitem(sys.modules, "market_data", market)
    monkeypatch.setattr(builtins, "open", forbidden)
    monkeypatch.setattr(Path, "open", forbidden)
    monkeypatch.setattr(Path, "write_text", forbidden)
    monkeypatch.setattr(sys, "argv", ["tool", "--unrecognized-option"])
    for name in MODULES:
        importlib.reload(modules[name])
    assert capsys.readouterr().out == ""


def test_risk_rules_and_flat_history(modules):
    history = modules["portfolio_history"]
    draw = modules["analyze_drawdown"]
    rows = [history.Row(f"2020-01-{i + 1:02}", r, a, 0, 1)
            for i, (r, a) in enumerate([(0, 100), (-0.1, 90), (0.1, 99), (0, 99)])]
    assert draw.sig_ma(rows, 2) == [True, True, False, True]
    assert draw.sig_lossstreak(rows, 2, 1) == [True, True, False, False]
    assert draw.apply_filter(rows, [True, False, True, False], 0.35) == pytest.approx([0, -0.035, 0.1, 0])
    assert modules["analyze_drawdown_bounded"].sig_count([1, 0.8, 0.8, 1], 0.9, 2) == [True, False, True, False]
    assert modules["analyze_drawdown_filters"].hysteresis([True, False, True, True], 1) == [True, False, False, True]
    assert math.isnan(history.mar(0.1, 0))
    flat = [history.Row(f"2020-01-{i + 1:02}", 0, 1, 0, 1) for i in range(12)]
    for name, args in [
        ("analyze_drawdown_liquidate", (0.9, 5, 10, 0.35, 0)),
        ("search_buy_block", (40, 0.9, 5, 10, 0.35)),
        ("search_max_profit", (40, 0.9, 5, 10)),
    ]:
        sig, rets = modules[name].simulate(flat, *args)
        assert all(sig) and rets == [0] * len(flat)